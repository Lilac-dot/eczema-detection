"""Revision Part 1: Title, Abstract, Introduction, Contributions."""
import docx

PATH = "Multi_Sensor_Wearable_AD_Monitoring_Paper_2026-09-15_revised.docx"
d = docx.Document(PATH)
paras = d.paragraphs

def set_text(idx, text):
    p = paras[idx]
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

# --- Title ---
set_text(0, "Development and Evaluation of a Multimodal Framework for Eczema Monitoring: "
             "Wearable Stress Detection, Image Classification, and Cross-Dataset Validation")

# --- Abstract ---
abstract = (
"Atopic dermatitis (AD), or eczema, is a chronic inflammatory skin condition whose "
"severity fluctuates in ways infrequent clinical visits cannot capture, motivating "
"interest in continuous, wearable- and image-based monitoring. This project investigates "
"whether independently validated wearable physiological and dermatological image "
"classification models can be assembled into a technically coherent multimodal "
"eczema-monitoring framework using only publicly available data, and identifies the "
"specific barriers currently preventing end-to-end validation of such a system. "
"Three components were developed and evaluated independently. A wearable stress "
"detector, trained on the public WESAD dataset (15 subjects, leave-one-subject-out "
"cross-validation), reached a mean AUC of 0.871 with LightGBM; per-subject inspection "
"revealed a severe decision-threshold instability concealed by this pooled figure, "
"which subject-specific baseline calibration substantially resolved (mean AUC 0.9405, "
"pooled F1 0.7602, mean per-subject F1 0.635), though one subject (S5) remained a "
"threshold-transfer failure. This validates general physiological stress detection on "
"healthy adults under a laboratory protocol, not an eczema-specific or patient-population "
"signal. An image-based eczema classifier was first trained on an Eczema-vs-Normal task "
"and reached 95.66% accuracy; auditing the data showed this result was driven by a "
"photographic-source shortcut (a four-fold brightness difference between classes) "
"rather than genuine lesion features, so the task was rebuilt from a same-source archive "
"comparing Eczema against seven clinically similar diseases, reaching a lower but "
"trustworthy internal result (81.07% accuracy, F1 81.18%). Zero-shot evaluation of this "
"model on two independent external datasets (SCIN, crowd-sourced consumer photography, "
"and SkinDisNet, a second clinical archive) showed near-chance discrimination (AUC "
"0.5347 and 0.4865), demonstrating that internal accuracy does not reflect cross-source "
"generalization. Generic data augmentation did not improve this; fine-tuning on real "
"external examples did, most effectively when the training sources were weighted "
"equally rather than in their natural proportion (SkinDisNet AUC rising to 0.6515), at "
"some cost to internal accuracy. A parallel attempt to detect sleep disruption from the "
"same class of wearable signal (AAUWSS dataset) produced no usable signal (mean AUC "
"0.4642, at chance), corroborated by two independent actigraphy baselines also scoring "
"at chance; this negative result is reported and the component is excluded from the "
"deployed pipeline. The wearable stress score and image classification score are "
"combined by an explicit, literature-motivated trigger-modulated rule rather than a "
"jointly trained model, since no public dataset pairs wearable recordings with skin "
"photographs from the same patients; this fusion stage is a working proof of concept "
"only, not fitted or validated against paired patient outcomes, and its output should "
"not be read as a calibrated probability of eczema severity or flare risk. The central "
"finding is methodological: each component can be independently developed, and its "
"failure modes characterized and partly addressed, using existing public data, but "
"validating the complete multimodal system requires a paired, longitudinal, "
"patient-level dataset linking wearable signals, skin photographs, and clinician-assessed "
"severity -- which does not currently exist in the public domain. A proposed low-cost "
"wearable hardware architecture, grounded in the models' actual sensor requirements, is "
"described but has not been built or tested."
)
set_text(8, abstract)

# --- Introduction ---
intro_p1 = (
"Atopic dermatitis is the most common inflammatory skin disease, affecting an estimated "
"15-20% of children and 1-3% of adults worldwide. Its severity is not static: it flares "
"and settles over days, driven by a mix of environmental, behavioural, and physiological "
"factors, including stress. Clinical practice still relies overwhelmingly on infrequent, "
"in-person visual assessment -- scoring systems such as SCORAD and EASI -- which are "
"episodic, subjective, and dependent on recall between visits. This creates an obvious "
"opening for continuous, objective monitoring using consumer or research-grade wearable "
"sensors and smartphone imaging, and a correspondingly large and fast-growing body of "
"literature attempting exactly that."
)
intro_p2 = (
"The research question this project addresses is: can independently validated wearable "
"physiological and image-based models be assembled into a technically coherent "
"multimodal eczema-monitoring framework using currently available public datasets, and "
"what limitations prevent validation of the complete system? To answer this, the project "
"investigates five things in turn: wearable stress detection, dermatological image "
"classification, cross-dataset generalization of that image classifier, decision-level "
"multimodal fusion, and the data requirements that would be needed for eventual "
"patient-level validation. The project's original goal was narrower -- a wearable system "
"estimating AD severity from a motion-derived nocturnal scratch signal combined with an "
"image-derived lesion score -- and evolved into this broader investigation in direct "
"response to evidence encountered during development: a hardware bandwidth limitation "
"documented in the literature (Section 2.3) ruled out the original scratch-detection "
"plan, and an image classifier's suspiciously high accuracy turned out to be a dataset "
"artifact rather than genuine diagnostic signal (Section 5.2). Rather than treating these "
"as detours, this report treats each as part of the experimental record: a case where a "
"plausible-looking result was checked rather than accepted."
)
intro_p3 = (
"The remainder of this report is organized as follows. Section 2 reviews the literature "
"this project is grounded in. Section 3 gives a system overview. Sections 4 through 6 "
"detail each stage -- wearable stress and sleep detection (Section 4), image-based "
"classification and its cross-dataset generalization (Section 5), and fusion (Section 6) "
"-- including what was tried, what failed, and why. Section 7 discusses the scientific "
"validity of the resulting system against the claims the project set out to test. "
"Section 8 sets out a proposed hardware architecture for an eventual wearable device. "
"Section 9 concludes and prioritizes future work. Section 10 records a later, direct "
"follow-up experiment (external validation of the image classifier) carried out after "
"the rest of the report was drafted."
)
set_text(11, intro_p1)
set_text(12, intro_p2)
set_text(13, intro_p3)

# --- Contributions ---
set_text(15,
    "This project's contributions are primarily methodological and experimental rather "
    "than clinical, and fall into five categories."
)

d.save(PATH)
print("Part 1 done: title, abstract, introduction, contributions intro saved.")
