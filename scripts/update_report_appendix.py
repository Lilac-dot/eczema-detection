"""
One-off script: inserts a comprehensive "Appendix A" reference table listing every
model trained in this project and its headline result, positioned right before the
References section (after the Section 10 postscript added by
update_report_postscript.py). Run once; not part of the regular pipeline.
"""
import docx

PATH = "Multi_Sensor_Wearable_AD_Monitoring_Paper_2026-09-15_revised.docx"

d = docx.Document(PATH)

refs_heading = None
for p in d.paragraphs:
    if p.text.strip() == "References" and p.style.name == "Heading 1":
        refs_heading = p
        break
assert refs_heading is not None, "Could not find References heading"


def add_para(text, style="Normal"):
    return refs_heading.insert_paragraph_before(text, style=style)


def add_table(rows, col_widths_hint=None):
    n_rows, n_cols = len(rows), len(rows[0])
    t = d.add_table(rows=n_rows, cols=n_cols, style="Table Grid")
    for r, row_vals in enumerate(rows):
        for c, val in enumerate(row_vals):
            cell = t.cell(r, c)
            cell.text = str(val)
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    refs_heading._p.addprevious(t._tbl)
    return t


add_para("Appendix A: Complete Model and Results Reference", style="Heading 1")
add_para(
    "Every model trained anywhere in this project, in the order it appears in the "
    "report, with its architecture/method, headline result, and current status "
    "(deployed / superseded / diagnostic-only / excluded). Deployed models are the "
    "ones the codebase actually calls at inference time; every other row is "
    "retained in the repository as part of the methodological record described in "
    "the relevant section, not as a live alternative."
)

rows = [
    ["Stage / Section", "Model", "Architecture / method", "Headline result", "Status"],

    ["A -- stress (4.4)", "LightGBM, pre-calibration", "Hand-crafted window stats, 26 features", "Mean AUC 0.871", "Superseded"],
    ["A -- stress (4.4)", "CNN+attention, pre-calibration", "4-branch 1D CNN, attention fusion", "Mean AUC 0.899; one fold AUC=1.0 with F1=0.00 (S2)", "Superseded"],
    ["A -- stress (4.6)", "LightGBM, post-calibration", "52 features (raw + subject-baseline-relative)", "Mean AUC 0.9405, pooled F1 0.760, mean per-subject F1 0.635, acc. 82.8%", "DEPLOYED"],
    ["A -- stress (4.6)", "CNN+attention, post-calibration", "Same architecture, calibrated inputs", "Mean AUC 0.889, F1 0.677; threshold range widened (S14 AUC=0.00)", "Comparison only"],
    ["A -- stress (4.7)", "CNN+attention, 3-seed ensemble", "3 independently-seeded copies, averaged", "Mean AUC 0.892, pooled AUC 0.886, F1 0.647", "Diagnostic, not deployed"],

    ["A-sleep (4.8)", "LightGBM (AAUWSS)", "Same recipe as Stage A-stress", "Mean LOSO AUC 0.4642, F1 0.1071, pooled AUC 0.3867 (below chance)", "Trained, excluded (negative result)"],
    ["A-sleep (4.8)", "Cole-Kripke actigraphy", "Fixed, non-trained formula (1992)", "Mean AUC 0.482, pooled 0.505", "Baseline comparison"],
    ["A-sleep (4.8)", "Sadeh actigraphy", "Fixed, non-trained formula (1994)", "Mean AUC 0.470, pooled 0.481", "Baseline comparison"],
    ["A-sleep (4.8)", "CNN+attention (AAUWSS)", "Same architecture as Stage A-stress", "Never completed -- CPU/memory limits on this hardware", "Inconclusive"],

    ["B -- image, original (5.1)", "ResNet18, Eczema vs. Normal", "CNN", "95.66% accuracy", "Superseded -- shortcut-inflated"],
    ["B -- image, original (5.1)", "LightGBM, Eczema vs. Normal", "90 colour features (15 colour spaces x mean+std)", "95.18% accuracy", "Superseded -- shortcut-inflated"],
    ["B -- image, original (5.2)", "LightGBM, brightness-normalized", "Same 90 features, normalized images", "93.98% accuracy", "Fix attempted, failed"],
    ["B -- image, original (5.2)", "ResNet18, cropped", "CNN, cropped images", "95.66% accuracy (unchanged)", "Fix attempted, failed"],
    ["B -- image, original (5.2)", "LightGBM, cropped", "Same 90 features, cropped images", "95.42% accuracy", "Fix attempted, failed"],

    ["B -- image, rebuilt (5.4)", "v1 CNN, frozen backbone", "ResNet18, only final layer trained", "69.55% accuracy, F1 46.74%", "Superseded"],
    ["B -- image, rebuilt (5.4)", "v1 LightGBM", "90 colour features", "82.43% accuracy, F1 57.06%", "Superseded"],
    ["B -- image, rebuilt (5.4)", "v2 CNN, fine-tuned", "ResNet18, layer4+fc unfrozen", "87.00% accuracy, F1 63.01%", "Superseded"],
    ["B -- image, rebuilt (5.4)", "Ensemble (v2 CNN + LightGBM)", "Averaged predictions", "89.05% accuracy, F1 67.84%", "Superseded"],
    ["B -- image, final (5.5)", "Final balanced CNN", "ResNet18, layer4+fc unfrozen, 50/50 balanced data", "81.07% accuracy, F1 81.18%, precision 79.92%, recall 82.47%", "DEPLOYED"],
    ["B -- image, diagnostic", "v3 colour-feature LightGBM", "90 colour features, same balanced data", "Internal: 70.41% accuracy, AUC 0.7829, F1 0.7036", "Diagnostic tool (shortcut checks)"],

    ["C -- fusion (6.2)", "fuse() -- flat additive rule", "Fixed formula, 0.5/0.5 weights (not trained)", "Runs correctly end-to-end; no accuracy metric, not validated against ground truth", "Retained for comparison only"],
    ["C -- fusion (6.2)", "trigger_index()/flare_risk() -- gated rule", "Fixed noisy-OR + multiplicative gate (not trained)", "Runs correctly end-to-end; no accuracy metric, not validated against ground truth", "Primary rule, proof-of-concept"],

    ["10.1 -- external validation", "Deployed CNN on internal test", "Same model as row 19, unmodified", "Acc. 0.8107, AUC 0.8644 (CI 0.8307-0.8973), F1 0.8118 (CI 0.7732-0.8493)", "Baseline reproduction"],
    ["10.1 -- external validation", "Deployed CNN on SCIN", "Zero-shot, no retraining", "Acc. 0.5092, AUC 0.5347 (CI 0.5003-0.5692), F1 0.1492 (CI 0.1119-0.1886)", "Negative result -- chance-level"],
    ["10.1 -- external validation", "Deployed CNN on SkinDisNet", "Zero-shot, no retraining", "Acc. 0.6673, AUC 0.4865 (CI 0.4559-0.5183), F1 0.1123 (CI 0.0794-0.1461)", "Negative result -- chance-level"],

    ["10.2 -- generalization", "Augmented CNN", "ResNet18, heavier augmentation, internal-only data", "Internal acc. 0.7515; SCIN AUC 0.5358; SkinDisNet AUC 0.5016", "No improvement"],
    ["10.2 -- generalization", "Multi-source CNN, natural mix", "ResNet18 fine-tuned on internal + SCIN-train + SkinDisNet-train (proportional)", "Internal AUC 0.8530/F1 0.7976; SCIN AUC 0.5783/F1 0.5214; SkinDisNet AUC 0.6090/F1 0.3823", "Real improvement (SkinDisNet)"],
    ["10.2 -- generalization", "Multi-source CNN, equal-weighted", "Same, but domain-equalized sampler", "Internal AUC 0.8317/F1 0.7758; SCIN AUC 0.5794/F1 0.5761; SkinDisNet AUC 0.6515/F1 0.4218", "Best external result achieved"],
]

add_table(rows)
add_para(
    "Table A1. Every model trained in this project, in report order, with its "
    "headline result and current status. \"DEPLOYED\" marks the two models the "
    "codebase actually calls at inference time (models/wesad_stress_lightgbm.txt "
    "for Stage A, models/curated_resnet18_balanced.pt for Stage B); every other row "
    "is retained in the repository for the record but not live."
)

d.save(PATH)
print("Saved.")
print("Appendix A inserted with", len(rows) - 1, "model rows.")
