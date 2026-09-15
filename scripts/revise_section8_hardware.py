"""Revision: Hardware section -- add implementation-status breakdown and
'Most Valuable Next Hardware Experiment' subsection; renumber existing 8.x."""
import docx

PATH = "Multi_Sensor_Wearable_AD_Monitoring_Paper_2026-09-15_revised.docx"
d = docx.Document(PATH)
paras = d.paragraphs

def find(text, style=None):
    for p in paras:
        if p.text.strip() == text and (style is None or p.style.name == style):
            return p
    return None

def set_text(p, text):
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

h81 = find("8.1 Why a single device does not work", "Heading 2")
h82 = find("8.2 Wrist sensor node", "Heading 2")
h83 = find("8.3 Hub", "Heading 2")
h84 = find("8.4 What this section does not claim", "Heading 2")
h85 = find("8.5 Future hardware: lesion-vs-surrounding-skin thermal imaging (not pursued)", "Heading 2")
assert all([h81, h82, h83, h84, h85])

set_text(h81, "8.2 Why a single device does not work")
set_text(h82, "8.3 Wrist sensor node")
set_text(h83, "8.4 Hub")
set_text(h84, "8.5 What this section does not claim")
set_text(h85, "8.7 Future hardware: lesion-vs-surrounding-skin thermal imaging (not the immediate next step)")

# Insert 8.1 Implementation status, right after h81's new position's original spot --
# i.e. before what is now "8.2 Why a single device does not work"
add_before = h81  # h81 now reads "8.2 ..." but the paragraph object is the same anchor
p_status_heading = add_before.insert_paragraph_before("8.1 Implementation status", style="Heading 2")
p_status_body = add_before.insert_paragraph_before(
    "This project's hardware involvement is entirely software and proposal at this stage. To keep "
    "that distinction explicit:"
)

def add_bullet(text):
    return add_before.insert_paragraph_before(text, style="List Bullet")

add_bullet("CURRENTLY IMPLEMENTED: the trained models themselves (Stage A-stress LightGBM, Stage B "
           "balanced CNN), their feature-extraction and preprocessing code, the inference functions "
           "(stage_a_stress_predict(), Stage B's classifier call), and the fusion code "
           "(trigger_index(), flare_risk(), fuse()) -- all real, working, and runnable today, on "
           "existing public data.")
add_bullet("PROPOSED, not built: the wrist sensor node (Section 8.3), its BLE streaming link, the "
           "Raspberry Pi 5 hub (Section 8.4), and camera integration for capturing Stage B input "
           "photographs on-device.")
add_bullet("NOT YET DONE: any physical hardware, any firmware, any sensor validation, any battery "
           "testing, any comparison against the Empatica E4 reference device the models were trained "
           "on, and any real-time deployment. No component has been purchased or fabricated.")

# --- Most Valuable Next Hardware Experiment, inserted before 8.7 (thermal imaging) ---
thermal_heading = find(
    "8.7 Future hardware: lesion-vs-surrounding-skin thermal imaging (not the immediate next step)",
    "Heading 2")
assert thermal_heading is not None

nh = thermal_heading.insert_paragraph_before("8.6 Most Valuable Next Hardware Experiment", style="Heading 2")
thermal_heading.insert_paragraph_before(
    "Given the implementation status in Section 8.1, the highest-value next hardware step is not a "
    "new capability but validating the one capability the deployed model actually depends on: build "
    "the wrist sensor node described in Section 8.3, and compare its raw signals -- EDA, "
    "temperature, PPG, acceleration -- directly against a worn Empatica E4 under a shared protocol. "
    "This matters because the deployed Stage A model was trained exclusively on E4 signal "
    "characteristics; whether a lower-cost sensor package reproduces those characteristics closely "
    "enough for the trained model to transfer is an open, testable, and currently untested question "
    "(Section 8.5). This experiment should precede, not follow, any attempt at patient-facing "
    "deployment: if the WESAD-trained model does not transfer to the new hardware's actual signal "
    "characteristics, no amount of further modelling work on the existing pipeline addresses that "
    "gap."
)

d.save(PATH)
print("Section 8 revised: implementation-status breakdown and Most Valuable Next Experiment added.")
